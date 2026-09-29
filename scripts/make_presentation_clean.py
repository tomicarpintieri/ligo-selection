"""Build clean 7-slide PPTX and PDF; fixed two-column grids prevent overlap."""
import sys, pathlib
ROOT=pathlib.Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'.ppt_build'))
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from matplotlib.backends.backend_pdf import PdfPages
import matplotlib.pyplot as plt
import matplotlib.image as mpimg

slides=[
('Sesgo de Selección y Observabilidad en LIGO',['El caso GW150914','Tomás Carpintieri · Agustín Miculicich','Universidad de Buenos Aires (UBA)','Proyecto final · 2026'],None),
('Motivación: el universo filtrado',['LIGO mide deformaciones diferenciales del espacio-tiempo.','Detectar requiere superar un umbral de SNR.','La detectabilidad depende de distancia, masa, orientación y posición celeste.','Población observada ≠ población intrínseca.'],None),
('Validación del pipeline: GW150914',['Control en datos reales antes de inferir selección.','SNR H1 = 19.810','SNR L1 = 13.543','Delay L1 − H1 = −7.080 ms','PSD y normalización verificadas con tests.'],'f01_snr_timeseries.png'),
('Waveform propia frente a LAL',['TaylorF2 no giratoria: amplitud 0PN y fase 3.5PN.','Match frente a referencia LAL = 1.000.','SNR inspiral = 11.358.','La inspiral sola subestima la detectabilidad de sistemas pesados.'],'f02_waveform_vs_lal.png'),
('Sensibilidad, horizonte y volumen',['ρ ∝ 1/D; por lo tanto V ∝ D³.','Horizonte óptimo de GW150914: 1969.9 Mpc.','Volumen euclídeo: 32.02 Gpc³.','En baja masa: Dₕ ∝ Mchirp^(5/6).'],'f03_horizon_vs_mchirp.png'),
('Geometría, red y puntos ciegos',['Un interferómetro L tiene cuatro direcciones ciegas exactas.','Promedio angular: ⟨F₊²⟩ ≈ ⟨F×²⟩ ≈ 1/5.','Horizonte/rango = 2.26487.','Tiempo máximo H1–L1 = 10.00274 ms.','El delay define un anillo compatible, no una localización puntual.'],None),
('Conclusiones',['El catálogo de detecciones es una población filtrada instrumentalmente.','El sesgo favorece fuentes cercanas, masivas y bien orientadas.','Pipeline reproducible: datos, resultados, procedencia, figuras y tests.','Siguiente paso: inyecciones, cosmología e inferencia jerárquica.'],None)]
navy=(13,24,43); white=RGBColor(245,247,250); teal=RGBColor(42,157,143)
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
for title,items,image in slides:
 s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(*navy)
 tb=s.shapes.add_textbox(Inches(.65),Inches(.45),Inches(12),Inches(.65)).text_frame.paragraphs[0];tb.text=title;tb.font.size=Pt(28);tb.font.bold=True;tb.font.color.rgb=white
 if title == slides[0][0]:
  tb.font.size=Pt(38);tb.alignment=1
  s.shapes[0].top=Inches(1.25);s.shapes[0].height=Inches(1.0)
 tf=s.shapes.add_textbox(Inches(.85),Inches(1.55),Inches(5.05 if image else 11.3),Inches(4.9)).text_frame;tf.clear()
 for i,item in enumerate(items):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text='• '+item;q.font.size=Pt(19);q.font.color.rgb=white;q.space_after=Pt(15)
 if title == slides[0][0]:
  tf.clear();tf.paragraphs[0].text='El caso GW150914';tf.paragraphs[0].alignment=1
  for item in items[1:]:q=tf.add_paragraph();q.text=item;q.alignment=1;q.font.size=Pt(20);q.font.color.rgb=white;q.space_after=Pt(12)
  tf.paragraphs[0].font.size=Pt(25);tf.paragraphs[0].font.color.rgb=teal
  s.shapes[1].left=Inches(1.0);s.shapes[1].top=Inches(3.0);s.shapes[1].width=Inches(11.3)
 if image:s.shapes.add_picture(str(ROOT/'figures'/image),Inches(6.45),Inches(1.5),width=Inches(6.15),height=Inches(4.5))
 # one understated accent only; no repeated footer text
 line=s.shapes.add_shape(1,Inches(.65),Inches(6.85),Inches(1.1),Inches(.05));line.fill.solid();line.fill.fore_color.rgb=teal;line.line.fill.background()
p.save(ROOT/'presentacion_sesgo_seleccion_ligo_limpia.pptx')

with PdfPages(ROOT/'presentacion_sesgo_seleccion_ligo_limpia.pdf') as pdf:
 for title,items,image in slides:
  fig=plt.figure(figsize=(13.333,7.5),facecolor='#0d182b');fig.text(.055,.9,title,color='white',fontsize=25,fontweight='bold')
  if title == slides[0][0]:
   fig.text(.5,.70,title,color='white',fontsize=34,fontweight='bold',ha='center');fig.text(.5,.43,'El caso GW150914',color='#2a9d8f',fontsize=24,ha='center');fig.text(.5,.30,'Tomás Carpintieri · Agustín Miculicich\nUniversidad de Buenos Aires (UBA) · Proyecto final · 2026',color='white',fontsize=19,ha='center');fig.texts[0].set_visible(False)
  ax=fig.add_axes([.065,.16,.38 if image else .84,.63]);ax.axis('off')
  for i,item in enumerate(items):ax.text(0,1-i*.17,'• '+item,color='white',fontsize=17,va='top',wrap=True)
  if image:
   iax=fig.add_axes([.51,.18,.44,.58]);iax.imshow(mpimg.imread(ROOT/'figures'/image));iax.axis('off')
  fig.add_artist(plt.Line2D([.055,.135],[.08,.08],color='#2a9d8f',lw=4,transform=fig.transFigure));pdf.savefig(fig);plt.close(fig)
