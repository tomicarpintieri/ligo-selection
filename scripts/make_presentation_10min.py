import sys, pathlib
ROOT=pathlib.Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'.ppt_build'))
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor

p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
navy=RGBColor(13,24,43);white=RGBColor(244,247,251);teal=RGBColor(42,157,143);muted=RGBColor(185,200,216)
def add(title, points=(), fig=None, note=''):
 s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=navy
 t=s.shapes.add_textbox(Inches(.7),Inches(.4),Inches(12),Inches(.7)).text_frame.paragraphs[0];t.text=title;t.font.size=Pt(29);t.font.bold=True;t.font.color.rgb=white
 tf=s.shapes.add_textbox(Inches(.85),Inches(1.45),Inches(5.4 if fig else 11.5),Inches(5.4)).text_frame;tf.clear()
 for i,x in enumerate(points):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=x;q.font.size=Pt(19);q.font.color.rgb=white;q.space_after=Pt(15)
 if fig:s.shapes.add_picture(str(ROOT/'figures'/fig),Inches(6.45),Inches(1.45),width=Inches(6.1))
 if note:
  n=s.notes_slide.notes_text_frame;n.text='Notas para ~1 min: '+note
 return s
add('Sesgo de selección y observabilidad en LIGO',['El caso GW150914','[Tu nombre] · UBA','Proyecto reproducible de ondas gravitacionales'],note='Presentar la pregunta: qué detecta el instrumento y qué queda sistemáticamente fuera.')
add('Mapa de la charla',['1. Control con GW150914','2. Señal, SNR y horizonte','3. Geometría: orientación y puntos ciegos','4. Red H1–L1 e implicancias'],note='Anticipar que el hilo conductor es pasar de una señal conocida a la función de selección.')
add('Motivación: el catálogo es un universo filtrado',['LIGO mide variaciones diferenciales de longitud con interferometría láser.','Detectar requiere SNR por encima de un umbral.','La probabilidad de detección depende de distancia, masa, orientación y posición celeste.','Por tanto: población observada ≠ población intrínseca.'],note='Explicar que el sesgo no es un detalle estadístico: cambia la demografía inferida.')
add('Control: recuperación de GW150914',['SNR H1 = 19.810','SNR L1 = 13.543','L1 llega 7.080 ms antes que H1.','El pipeline recupera una detección ya establecida antes de hacer inferencias nuevas.'],'f01_snr_timeseries.png',note='Leer picos y explicar por qué este control hace confiable lo que sigue.')
add('Waveform propia: contraste con LAL',['TaylorF2 no giratoria, fase 3.5PN y amplitud 0PN.','Match contra la referencia LAL = 1.000.','SNR inspiral = 11.358: menor que IMR porque GW150914 es pesado.','Consecuencia: inspiral-only es una cota inferior de detectabilidad.'],'f02_waveform_vs_lal.png',note='Distinguir validación de forma de onda de alcance físico de la aproximación.')
add('De SNR a alcance',['ρ ∝ 1/D y V ∝ D³.','Con ρ*=8: horizonte óptimo = 1969.9 Mpc.','Volumen euclídeo = 32.02 Gpc³.','En baja masa: D_h ∝ Mchirp^(5/6).'],'f03_horizon_vs_mchirp.png',note='Destacar que un pequeño cambio en alcance se amplifica cúbicamente en volumen.')
add('Geometría: por qué el horizonte no es el rango',['Un interferómetro L tiene cuatro direcciones ciegas exactas.','<F+²> ≈ <Fx²> ≈ 1/5 para un cielo isotrópico.','Factor horizonte/rango = 2.26487.','El rango incorpora orientación, polarización y distribución en el cielo.'],note='Aclarar que el horizonte es el mejor caso; el rango es la cantidad útil para conteos.')
add('Red H1–L1',['ρ_red² = ρ_H1² + ρ_L1².','Tiempo máximo de viaje: 10.00274 ms.','Ganancia media de red = 2.0018 sobre un detector.','El delay determina un anillo compatible; no una localización puntual.'],note='Explicar cobertura complementaria y por qué una red mayor mejora localización.')
add('Rigor técnico y reproducibilidad',['Datos hasheados y PSD validada contra referencia.','Resultados en results.json; números y figuras con provenance.','run_all.py reconstruye el análisis.','39 tests automatizados aprobados.'],note='Usar esta placa para defender el método, no solo los números.')
add('Conclusiones',['No se puede hacer demografía de agujeros negros sin modelar selección.','El sesgo favorece fuentes cercanas, masivas y bien orientadas.','El pipeline establece una base para campañas de inyección.','Siguiente paso: cosmología, red ampliada e inferencia jerárquica.'],note='Cerrar volviendo a la tesis: el instrumento filtra el universo observado.')
p.save(ROOT/'presentacion_sesgo_seleccion_ligo_10min.pptx')
