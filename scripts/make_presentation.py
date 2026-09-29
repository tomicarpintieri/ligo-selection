import sys
import pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / '.ppt_build'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / 'presentacion_ligo_selection.pptx'
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
bg=(15,23,42); white=RGBColor(242,245,249); teal=RGBColor(42,157,143); orange=RGBColor(232,112,58)

def slide(title, bullets=(), image=None, subtitle=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(*bg)
    box=s.shapes.add_textbox(Inches(.7),Inches(.45),Inches(12),Inches(.7)).text_frame
    p=box.paragraphs[0];p.text=title;p.font.size=Pt(30);p.font.bold=True;p.font.color.rgb=white
    if subtitle:
        q=box.add_paragraph();q.text=subtitle;q.font.size=Pt(14);q.font.color.rgb=RGBColor(170,190,210)
    if bullets:
        tf=s.shapes.add_textbox(Inches(.9),Inches(1.5),Inches(5.4 if image else 11.5),Inches(5.3)).text_frame;tf.clear()
        for i,b in enumerate(bullets):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=b;p.font.size=Pt(21);p.font.color.rgb=white;p.space_after=Pt(16)
    if image: s.shapes.add_picture(str(ROOT/'figures'/image),Inches(6.5),Inches(1.45),width=Inches(6.1))
    return s

slide('LIGO Selection', ['Qué detecta LIGO, qué deja afuera y por qué la selección altera el catálogo observado.'], subtitle='Proyecto reproducible sobre GW150914')
slide('La pregunta física', ['Las fusiones no se observan de forma uniforme.', 'La detectabilidad depende de distancia, masa, orientación, posición en el cielo y red de detectores.', 'Objetivo: medir ese sesgo con un pipeline verificable.'])
slide('Primero: un control conocido', ['El pipeline recupera GW150914 en datos reales.', 'SNR Hanford: 19.810', 'SNR Livingston: 13.543', 'Llegada de L1: 7.080 ms antes que H1.'], 'f01_snr_timeseries.png')
slide('La señal: TaylorF2 frente a LAL', ['Construimos una waveform inspiral en frecuencia.', 'Relación de amplitud mediana: 1.000.', 'Match fase/tiempo: 1.000.', 'Limitación: una inspiral sola subestima sistemas pesados; faltan merger y ringdown.'], 'f02_waveform_vs_lal.png')
slide('De SNR a distancia', ['SNR ∝ 1/distancia.', 'Con umbral SNR = 8: horizonte GW150914 = 1969.9 Mpc.', 'Volumen euclídeo sensible = 32.02 Gpc³.', 'La escala de baja masa sigue Mchirp^(5/6).'], 'f03_horizon_vs_mchirp.png')
slide('Orientación: el detector no oye igual todo el cielo', ['El patrón de antena tiene cuatro direcciones ciegas exactas.', '<F+²> = 0.19997 y <Fx²> = 0.19987, consistentes con 1/5.', 'El factor de proyección RMS es 0.39996.', 'Horizonte / rango = 2.26487: el horizonte es el mejor caso, no el promedio.'])
slide('La red H1–L1', ['La separación máxima medida por tiempo de vuelo es 10.00274 ms.', 'La ganancia promedio de respuesta de la red es 2.0018 frente a un detector.', 'El delay de GW150914 es consistente con un anillo en el cielo, no con una localización puntual.'])
slide('Qué aprendimos', ['Los catálogos favorecen fuentes más masivas, cercanas y bien orientadas.', 'Una detección no es una muestra representativa del universo.', 'La selección debe modelarse antes de inferir poblaciones astrofísicas.'])
slide('Rigor y reproducibilidad', ['Resultados en results.json; cada cifra tiene procedencia.', 'Scripts por etapa y figuras regenerables.', 'run_all.py reconstruye el análisis.', '39 tests automatizados aprobados.'])
slide('Limitaciones y próximos pasos', ['Volumen euclídeo: sin cosmología ni corrimiento al rojo.', 'La red actual aún no incluye mapas/animación finales de cobertura.', 'Próximo paso: campaña de inyecciones con población de masas, orientación y distancia.'])
slide('Cierre', ['La pregunta no es solo “¿se detectó?”', 'También es “¿qué tipo de universo permite detectar este instrumento?”'], subtitle='Gracias')
prs.save(OUT)
print(OUT)
