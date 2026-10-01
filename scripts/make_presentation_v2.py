"""Build an alternative five-minute deck: the same talk, led by the figures.

This does NOT replace scripts/make_presentation_5min.py. Both scripts and both
decks are kept so the two directions can be compared side by side before one is
chosen.

What differs from v1:

  ground      light instead of navy, so the matplotlib figures -- which are
              rendered on white -- sit on the page instead of floating as white
              panels on a dark field
  palette     taken from src/gwsel/figures.py, so the slides and the plots are
              one visual system rather than two
  figures     four instead of one. The antenna pattern now carries the main
              result, which in v1 was numbers alone, and the sidereal sweep gets
              a slide because it is the mechanism behind that result
  structure   eight slides instead of seven; the four the brief asks for are
              unchanged and still map one to one

The content is deliberately the same. Nothing was added to the argument.

Run:
    py -3.12 -m venv .venv-presentation
    .venv-presentation\\Scripts\\python.exe -m pip install -r requirements-presentation.lock
    .venv-presentation\\Scripts\\python.exe scripts\\make_presentation_v2.py
"""

import json
from pathlib import Path

from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentacion_final_ligo_5min_v2.pptx"
BUILD = ROOT / ".ppt_build"
R = json.loads((ROOT / "results.json").read_text(encoding="utf-8"))

# Palette lifted from src/gwsel/figures.py so the deck and the plots agree.
GROUND = RGBColor(0xF7, 0xF8, 0xF9)
PALE = RGBColor(0xEE, 0xF1, 0xF4)
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x6B, 0x76, 0x84)
RULE = RGBColor(0xDF, 0xE3, 0xE8)
SIGNAL = RGBColor(0xD1, 0x59, 0x2A)      # the plots' orange, darkened for paper
TEAL = RGBColor(0x1F, 0x7F, 0x72)
DEEP = RGBColor(0x16, 0x22, 0x2E)
ON_DEEP = RGBColor(0xE4, 0xEA, 0xEF)

DISPLAY = "Georgia"      # titles only
BODY = "Aptos"           # everything else
TECH = "Consolas"        # eyebrows, units, small annotations

W, H = 13.333, 7.5


# --------------------------------------------------------------------- helpers
def text(slide, s, x, y, w, h, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=BODY, spacing=None, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(s.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        f = p.font
        f.name, f.size, f.bold, f.color.rgb = font, Pt(size), bold, color
    return box


def band(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.shadow.inherit = False
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    return shape


def picture(slide, name, x, y, width=None, height=None):
    """Place a figure, white margins trimmed. Aspect is never distorted."""
    path = _trimmed(name)
    im = Image.open(path)
    aspect = im.size[0] / im.size[1]
    if width is None:
        width = height * aspect
    if height is None:
        height = width / aspect
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                    Inches(width), Inches(height))


def _trimmed(name):
    """Drop the uniform white border matplotlib leaves. Nothing else is cut.

    Cropping the plot itself was tried and abandoned: a pptx cannot be previewed
    from here, so a crop that clipped an axis would have shipped unseen. The
    English figure titles therefore stay; each slide carries its own Spanish
    headline above them.
    """
    BUILD.mkdir(exist_ok=True)
    out = BUILD / f"trim_{name}"
    src = ROOT / "figures" / name
    if out.exists() and out.stat().st_mtime > src.stat().st_mtime:
        return out
    im = Image.open(src).convert("RGB")
    box = ImageChops.difference(im, Image.new("RGB", im.size, (255, 255, 255))).getbbox()
    im.crop(box).save(out) if box else im.save(out)
    return out


def note(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def page(prs, eyebrow, title, subtitle=None, dark=False):
    """A content slide: eyebrow, rule, title, optional standfirst."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DEEP if dark else GROUND
    ink = ON_DEEP if dark else INK
    text(slide, eyebrow.upper(), .78, .5, 8.0, .25, 10.5, SIGNAL, True, font=TECH)
    band(slide, .78, .84, 1.5, .035, SIGNAL)
    # 0.95 of height, not 0.72: a two-line title ran into the standfirst below.
    text(slide, title, .78, 1.06, 11.7, .95, 28, ink, True, font=DISPLAY, spacing=1.08)
    if subtitle:
        text(slide, subtitle, .78, 2.04, 10.6, .4, 14.5, MUTED)
    text(slide, "ligo-selection  ·  proyecto final", .78, 7.02, 5.0, .22, 9, MUTED, font=TECH)
    return slide


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)

# ---------------------------------------------------------------- 1 · portada
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = GROUND
text(s, "PROYECTO FINAL  ·  5 MINUTOS", .9, 1.15, 6.0, .3, 11, SIGNAL, True, font=TECH)
band(s, .9, 1.58, 1.9, .045, SIGNAL)
text(s, "¿De qué es ciego\nLIGO?", .9, 1.95, 6.4, 2.1, 46, INK, True, font=DISPLAY, spacing=.95)
text(s, "Efectos de selección en el catálogo de fusiones",
     .92, 4.25, 6.1, .45, 19, MUTED)
band(s, .92, 5.12, 6.1, .012, RULE)
text(s, "Tomás Carpintieri  ·  Agustín Miculicich", .92, 5.42, 6.2, .32, 16, INK, True)
text(s, "Ondas Gravitacionales e Investigación Asistida por IA\nUniversidad de Buenos Aires",
     .92, 5.82, 6.2, .68, 13, MUTED, spacing=1.25)
picture(s, "f04_antenna_pattern.png", 7.45, 2.05, width=5.3)
text(s, "respuesta del detector: oscuro = sordo", 7.45, 5.52, 5.3, .25, 10.5, MUTED,
     align=PP_ALIGN.CENTER, font=TECH)
note(s, "20 s. Una detección no es una muestra neutral del universo: LIGO oye "
        "mejor unas fusiones que otras. Esa imagen de la derecha es la respuesta "
        "del propio detector, y es de donde sale todo lo demás.")

# ---------------------------------------------------------------- 2 · problema
s = page(prs, "1 · Problema", "El catálogo observado no es el universo",
         "La detectabilidad depende tanto de la fuente como de la geometría del detector.")
cols = [(1.0, "LA FUENTE", "masa\ndistancia"),
        (5.15, "LA GEOMETRÍA", "posición en el cielo\norientación"),
        (9.3, "EL INSTRUMENTO", "ruido\numbral de detección")]
for x, head, body in cols:
    band(s, x, 2.72, 3.0, .035, INK)
    text(s, head, x, 2.95, 3.0, .28, 11, SIGNAL, True, font=TECH)
    text(s, body, x, 3.38, 3.0, .95, 21, INK, True, spacing=1.2)
for x in (4.35, 8.5):
    text(s, "→", x, 3.42, .5, .45, 26, MUTED, True, align=PP_ALIGN.CENTER)
band(s, 1.0, 5.25, 11.3, 1.0, DEEP)
text(s, "Se detectan preferentemente fuentes cercanas, masivas y bien orientadas.",
     1.3, 5.56, 10.7, .4, 20, ON_DEEP, True, align=PP_ALIGN.CENTER, font=DISPLAY)
note(s, "45 s. La pregunta no es si una fusión ocurrió, sino qué probabilidad "
        "había de que esta red la detectara. Esa probabilidad sesga cualquier "
        "conclusión sobre poblaciones de agujeros negros.")

# ------------------------------------------------------- 3 · resultado principal
s = page(prs, "2 · Resultado principal",
         "El horizonte óptimo sobrestima el alcance típico",
         "Promediando cielo, orientación y polarización, la fuente típica se ve mucho más cerca.")
picture(s, "f04_antenna_pattern.png", .78, 2.58, width=6.3)
text(s, "Los cuatro cruces son ceros exactos: direcciones en las que el detector\n"
        "no mide nada, por más fuerte que sea la onda.",
     .78, 6.3, 6.3, .55, 11.5, MUTED, spacing=1.3)
text(s, "÷ 2.2649", 8.1, 2.62, 4.4, .95, 52, SIGNAL, True, font=DISPLAY)
text(s, "horizonte  /  rango", 8.15, 3.62, 4.4, .3, 13, MUTED, font=TECH)
band(s, 8.1, 4.12, 4.4, .012, RULE)
for i, (value, unit, label) in enumerate([
        (f"{R['horizon_gw150914_mpc']:,.0f}".replace(",", " "), "Mpc",
         "horizonte óptimo de GW150914"),
        (f"{R['sensitive_volume_gpc3']:.1f}", "Gpc³", "volumen euclídeo accesible")]):
    y = 4.42 + i * 1.08
    text(s, value, 8.1, y, 2.0, .5, 30, INK, True, font=DISPLAY)
    text(s, unit, 10.2, y + .17, 1.0, .3, 14, MUTED, font=TECH)
    text(s, label, 8.12, y + .55, 4.4, .28, 11.5, MUTED)
text(s, "El volumen cambia como la distancia al cubo: perder alcance se paga caro.",
     8.1, 6.42, 4.45, .45, 12, INK, True, spacing=1.25)
note(s, "55 s. Éste es el resultado central. El horizonte es el mejor caso "
        "posible: fuente justo arriba y bien orientada. Al promediar sobre todo "
        "el cielo y todas las orientaciones, el alcance equivalente en volumen "
        "es 2.2649 veces menor. Como el volumen va con la distancia al cubo, "
        "ignorar la geometría cambia mucho los conteos.")

# -------------------------------------------------------- 4 · el cielo que rota
s = page(prs, "Por qué · geometría", "Y los puntos ciegos barren el cielo cada día",
         "El patrón está atornillado a la Tierra, así que gira con ella.")
picture(s, "f06_rotation_panels.png", 2.9, 2.48, height=3.95)
text(s, "Una misma fuente puede ser audible o invisible según la hora a la que ocurra.",
     1.0, 6.62, 11.3, .3, 14, INK, True, align=PP_ALIGN.CENTER)
note(s, "30 s. Esto es lo que hay detrás del factor anterior. El mapa de "
        "sensibilidad está pegado al detector, y la Tierra gira: las zonas "
        "oscuras se mueven. Una fusión que a una hora caía en un punto ciego, "
        "seis horas después se escucha perfectamente.")

# --------------------------------------------------------------- 5 · validación
s = page(prs, "3 · Validación", "Antes de estimar selección, recuperamos GW150914",
         "El mismo código encuentra la señal conocida en datos reales de los dos detectores.")
picture(s, "f01_snr_timeseries.png", .78, 2.62, width=7.5)
for i, (value, label) in enumerate([
        (f"{R['h1_peak_snr']:.2f}", "SNR en Hanford"),
        (f"{R['l1_peak_snr']:.2f}", "SNR en Livingston"),
        (f"{R['time_delay_l1_minus_h1_ms']:.2f} ms", "L1 llega antes que H1")]):
    y = 2.72 + i * 1.22
    band(s, 8.75, y, 3.8, .95, PALE)
    text(s, value, 9.05, y + .17, 1.9, .5, 25, INK, True, font=DISPLAY)
    text(s, label, 9.05, y + .62, 3.3, .25, 11.5, MUTED)
text(s, f"Contra la referencia de LAL, match = {R['match_vs_lal']:.3f}",
     8.75, 6.5, 3.9, .3, 11.5, INK)
text(s, f"Máximo retardo geométrico H1–L1 = {R['h1_l1_light_travel_ms']:.3f} ms",
     8.75, 6.84, 3.9, .25, 9.5, MUTED, font=TECH)
note(s, "55 s. Antes de afirmar nada nuevo, el pipeline tiene que encontrar algo "
        "que ya se conoce. Recupera SNR 19.81 en Hanford, 13.54 en Livingston y "
        "el retraso de 7.08 ms. La geometría de los sitios da un máximo de "
        "10.013 ms, así que ese retraso no da una posición: da un anillo.")

# --------------------------------------------------------------- 6 · aprendizaje
s = page(prs, "4 · Experiencia / aprendizaje",
         "Los tests que parecían de más encontraron los errores",
         "Cada número quedó atado a un dato, una prueba y una limitación declarada.")
items = [(1.0, "01", "Un test borrado",
          "El corte de la forma de onda estaba desactivado. El test que lo "
          "detectaba había sido eliminado del repositorio."),
         (5.15, "02", "No movimos el margen",
          "La varianza de blanqueo daba 0.442 donde debía dar 0.965. Buscamos "
          "la causa en vez de ensanchar la tolerancia."),
         (9.3, "03", "Procedencia de todo",
          "Cada número y cada figura registran qué los produjo y contra qué se "
          "verificaron. Un test lo exige.")]
for x, n, head, body in items:
    band(s, x, 2.72, 3.0, 3.1, PALE)
    text(s, n, x + .28, 2.95, .8, .42, 19, SIGNAL, True, font=TECH)
    text(s, head, x + .28, 3.46, 2.5, .5, 17, INK, True, font=DISPLAY)
    text(s, body, x + .28, 4.08, 2.5, 1.6, 12.5, INK, spacing=1.32)
band(s, 1.0, 6.18, 11.3, .72, DEEP)
text(s, "Un resultado confiable no es un número: es un número que se puede volver a producir y discutir.",
     1.3, 6.37, 10.7, .35, 15, ON_DEEP, True, align=PP_ALIGN.CENTER)
note(s, "45 s. Ésta fue la dificultad real. Un test reveló que el corte de la "
        "forma de onda estaba apagado, y ese test había sido borrado del "
        "repositorio: los chequeos que parecen redundantes cuando los escribís "
        "son los que encuentran algo cuando los corrés. Y cuando el blanqueo dio "
        "0.442 en vez de 0.965 no tocamos la tolerancia: buscamos la causa y "
        "llegamos a 0.959.")

# -------------------------------------------------------------------- 7 · cierre
s = page(prs, "Cierre", "Modelar la selección es el primer paso", dark=True)
text(s, "El universo que ve LIGO es una población filtrada por distancia, masa y geometría.",
     .78, 2.5, 11.5, .55, 23, ON_DEEP, True, font=DISPLAY, spacing=1.2)
band(s, .78, 3.42, 11.5, .012, RGBColor(0x33, 0x44, 0x55))
text(s, "EL PILOTO, EN RUIDO REAL", .78, 3.78, 5.4, .25, 10.5, SIGNAL, True, font=TECH)
text(s, f"De {R['injection_count']:.0f} fusiones inyectadas en datos reales de H1 y L1 "
        f"recuperamos el {R['injection_recovery_fraction']:.1%}, frente a un "
        f"{R['injection_background_false_alarm_fraction']:.1%} de falsas alarmas "
        f"al mismo umbral.",
     .78, 4.18, 5.6, 1.3, 15, ON_DEEP, spacing=1.3)
text(s, "SIGUIENTE PASO", 7.1, 3.78, 5.2, .25, 10.5, SIGNAL, True, font=TECH)
text(s, "Una campaña completa: incluir fusión y ringdown, ampliar la población y "
        "medir la función de selección de la red.",
     7.1, 4.18, 5.2, 1.3, 15, ON_DEEP, spacing=1.3)
text(s, "Gracias", .78, 6.05, 4.0, .5, 24, SIGNAL, True, font=DISPLAY)
note(s, "30 s. Cerramos con lo medido: el catálogo favorece fuentes cercanas, "
        "bien orientadas y, en el régimen de inspiral, más masivas. El piloto en "
        "ruido real recuperó 39.6 % frente a 4.2 % de fondo al mismo umbral. El "
        "paso siguiente es la campaña completa con fusión y ringdown.")

# ------------------------------------------------------------------- 8 · gracias
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = GROUND
text(s, "Gracias por escucharnos", .9, .78, 11.5, .7, 34, INK, True,
     font=DISPLAY, align=PP_ALIGN.CENTER)
text(s, "Esperamos sus preguntas", .9, 1.52, 11.5, .35, 16, MUTED, align=PP_ALIGN.CENTER)
# The figure carries its own caption; a second one landed on top of it.
picture(s, "f05_network_skymap.png", 1.05, 2.22, width=11.2)
band(s, 5.32, 6.44, 2.7, .012, SIGNAL)
text(s, "Tomás Carpintieri  ·  Agustín Miculicich      Universidad de Buenos Aires",
     .9, 6.7, 11.5, .3, 13, MUTED, align=PP_ALIGN.CENTER)
note(s, "Dejar visible durante las preguntas. Si alguien pregunta por la "
        "localización, el mapa está a la vista: el retraso define un anillo, no "
        "un punto.")

prs.save(OUT)
print(OUT)
